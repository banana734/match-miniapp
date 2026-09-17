# 从「匹配信息.xlsx」里抽两侧各 10 条最完整的资料，输出成 JSON 供 Node 录入数据库。
# 用法：D:/anaconda/ana/python.exe scripts/extract-seed-data.py
import json
import re
import openpyxl

SRC = r"D:/桌面/文件/友导师大创/匹配信息.xlsx"
OUT = r"D:/VScode/Match/scripts/seed-data.json"
PICK = 10

# 和 app 的选项表保持一致（src/constants/profile-options.js）
SUBJECT_OPTIONS = ['语文', '数学', '英语', '物理', '化学', '生物', '政治', '历史',
                   '体育', '绘画', '音乐', '舞蹈', '地理', '其他']

# 这些是「其他___」填空列里常见的无意义填法，一律当没填
PLACEHOLDERS = {'无', '无。', '没有', '暂无', '/', '-', 'n/a', 'N/A', '空', '无 '}

FAMILY_COLUMNS = [
    (r'1、孩子姓名', 'name', 'text'),
    (r'^2、孩子性别', 'gender', 'text'),
    (r'^3、孩子当前年级', 'grade', 'text'),
    (r'4、\(1\)家长称呼', 'parentName', 'text'),
    (r'4、\(2\)', 'phone', 'text'),
    (r'^微信号$', 'wechat', 'text'),
    (r'^5、家庭常住区域', 'area', 'text'),
    (r'^其他省市', 'areaOther', 'text'),
    (r'^6、孩子目前最需要辅导的科目', 'subjects', 'list'),
    (r'^其他\(可以填：无\)$', 'subjectOther', 'text'),
    (r'^7、您认为孩子最主要的困难类型', 'difficulties', 'list'),
    (r'^其他困难', 'difficultyOther', 'text'),
    (r'^8、孩子以前是否参加过课外辅导', 'intro', 'text'),
    (r'^9、您认为一位优秀的家教老师', 'teacherTraits', 'list'),
    (r'^10、在辅导过程中', 'teachingStyles', 'list'),
    (r'^11、您希望家教老师在每次辅导时', 'mainFocus', 'text'),
    (r'^其他方面', 'mainFocusOther', 'text'),
    (r'^12、除了辅导功课', 'learningState', 'text'),
    (r'^13、对于家教老师的沟通', 'communicationExpectation', 'text'),
    (r'^其他期待', 'communicationExpectationOther', 'text'),
    (r'^14、您是否理解并认同', 'understanding', 'text'),
    (r'^15、为持续优化项目', 'feedbackWillingness', 'text'),
    (r'^16、您还有什么想特别说明', 'extraNote', 'text'),
    (r'^意向上课方式', 'classModes', 'list'),
    (r'^上课频率', 'classFrequency', 'text'),
]

MENTOR_COLUMNS = [
    (r'^姓名$', 'name', 'text'),
    (r'^性别$', 'gender', 'text'),
    (r'^是否参加友导师项目$', 'mentorProject', 'text'),
    (r'^骨干成员$', 'coreMember', 'text'),
    (r'^学校$', 'school', 'text'),
    (r'^学院$', 'college', 'text'),
    (r'^专业$', 'major', 'text'),
    (r'^年级$', 'grade', 'text'),
    (r'^电话联系方式-手机号$', 'phone', 'text'),
    (r'^微信号$', 'wechat', 'text'),
    (r'^擅长科目$', 'mentorSubjects', 'list'),
    (r'^其他擅长科目', 'mentorSubjectOther', 'text'),
    (r'^意向教学年级段$', 'mentorTeachingGradeRange', 'text'),
    (r'^风格类型$', 'mentorStyleTypes', 'list'),
    (r'^意向上课方式$', 'mentorTeachingModes', 'list'),
    (r'^暑假所在地$', 'mentorSummerLocation', 'text'),
    (r'^开学后所在地$', 'mentorSchoolLocation', 'text'),
    (r'^上课频率$', 'mentorClassFrequency', 'text'),
]

NORMALIZE_REPORT = []


def clean(value):
    if value is None:
        return ''
    return re.sub(r'\s+', ' ', str(value)).strip()


def is_placeholder(text):
    return text in PLACEHOLDERS


def normalize_subject(raw):
    """科目归一化：'语文/作文' -> '语文'；其余原样；不在选项表里的记下来报警。"""
    text = raw.strip()
    if '/' in text:
        first = text.split('/')[0].strip()
        if first in SUBJECT_OPTIONS:
            NORMALIZE_REPORT.append('%s -> %s' % (text, first))
            return first
    if text and text not in SUBJECT_OPTIONS:
        NORMALIZE_REPORT.append('%s（不在选项表）' % text)
    return text


def to_list(text):
    if not text:
        return []
    parts = [p.strip() for p in text.split(',')]
    return [p for p in parts if p]


def extract(sheet_name, columns):
    ws = wb[sheet_name]
    headers = [clean(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

    # 先按正则把「字段 -> 列下标」定下来
    index_of = {}
    for pattern, field, kind in columns:
        for index, header in enumerate(headers):
            if re.search(pattern, header):
                index_of[field] = (index, kind)
                break

    rows = []
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        record = {}
        filled = 0
        for field, (index, kind) in index_of.items():
            text = clean(row[index]) if index < len(row) else ''

            if kind == 'list':
                if field in ('subjects', 'mentorSubjects'):
                    items = [normalize_subject(v) for v in to_list(text)]
                    items = [v for v in items if v]
                else:
                    items = to_list(text)
                record[field] = items
                if items:
                    filled += 1
                continue

            if text and is_placeholder(text):
                text = ''
            record[field] = text
            if text:
                filled += 1

        if not record.get('name'):
            continue

        rows.append({'rowNumber': row_number, 'filled': filled, 'profile': record})

    rows.sort(key=lambda item: (-item['filled'], item['rowNumber']))
    return rows[:PICK]


wb = openpyxl.load_workbook(SRC, data_only=True)

family_rows = extract('家庭', FAMILY_COLUMNS)
mentor_rows = extract('友导师', MENTOR_COLUMNS)

print('归一化记录：')
for line in sorted(set(NORMALIZE_REPORT)):
    print('   ', line)
print()

for label, rows in (('家庭', family_rows), ('友导师', mentor_rows)):
    print('%s 取中 %d 条：' % (label, len(rows)))
    for item in rows:
        profile = item['profile']
        subtitle = profile.get('area') or (str(profile.get('school', '')) + ' · ' + str(profile.get('major', '')))
        print('    第%3d行  完整度%2d  %-8s  %s' % (
            item['rowNumber'], item['filled'], profile['name'], subtitle))
    print()

payload = {'family': family_rows, 'mentor': mentor_rows}
with open(OUT, 'w', encoding='utf-8') as handle:
    json.dump(payload, handle, ensure_ascii=False, indent=2)
print('已写出：' + OUT)
